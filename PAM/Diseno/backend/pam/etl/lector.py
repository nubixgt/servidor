"""Lectura en streaming de la matriz PAM desde CSV o Excel (.xlsx).

- Detecta codificación (UTF-8, Mac Roman, Windows-1252/Latin-1) y separador (; , tab |).
- Reconoce encabezados por nombre aunque cambien mayúsculas, tildes o espacios,
  y encuentra la fila de encabezados en las primeras 15 filas.
- Entrega filas como listas de texto en el orden de COLUMNAS (columnas ausentes = None).
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
from pathlib import Path
from typing import Iterator

# columna destino en stg.matriz -> nombres aceptados en el archivo (normalizados)
COLUMNAS: list[tuple[str, list[str]]] = [
    ("departamento", ["departamento", "depto"]),
    ("municipio", ["municipio"]),
    ("no_expediente", ["no. expediente", "no expediente", "numero expediente", "expediente", "no. de expediente"]),
    ("nombre", ["nombre", "nombre completo"]),
    ("dpi", ["dpi", "cui", "dpi/cui", "no. dpi"]),
    ("fecha_nacimiento", ["fecha nacimiento", "fecha de nacimiento"]),
    ("fecha_solicitud", ["fecha solicitud", "fecha de solicitud"]),
    ("anio_solicitud", ["ano de solicitud", "anio de solicitud", "ano solicitud"]),
    ("telefono", ["telefono", "tel"]),
    ("direccion", ["direccion"]),
    ("fecha_observacion", ["fecha observacion", "fecha de observacion"]),
    ("ultima_observacion", ["ultima observacion general", "ultima observacion", "observacion"]),
    ("id_apoderado", ["id apoderado", "tipo id apoderado"]),
    ("no_id_apoderado", ["no. id apoderado", "no id apoderado", "numero id apoderado"]),
    ("apoderado", ["apoderado", "nombre apoderado"]),
    ("visita_positiva", ["visita positiva", "fecha visita positiva"]),
    ("anio_visita_positiva", ["ano de visita positiva"]),
    ("mes_visita", ["mes de visita", "mes de visita positiva"]),
    ("visita_negativa", ["visita negativa", "fecha visita negativa"]),
    ("anio_visita_negativa", ["ano de visita negativa"]),
    ("mes_visita_negativa", ["mes de visita negativa"]),
    ("nomina_desde", ["en nomina desde", "nomina desde"]),
    ("anio_ingreso_nomina", ["ano de ingreso a nomina"]),
    ("nomina_hasta", ["en nomina hasta", "nomina hasta"]),
    ("fallecio", ["fallecio", "fecha fallecimiento", "fecha de fallecimiento"]),
    ("anio_fallecido", ["ano de fallecido"]),
    ("suspendido", ["suspendido", "fecha suspension"]),
    ("fecha_ultima_resolucion", ["fecha ultima resolucion", "fecha de ultima resolucion"]),
    ("anio_resolucion", ["ano de resolucion"]),
    ("tipo_resolucion", ["tipo resolucion", "tipo de resolucion"]),
    ("resolucion", ["resolucion"]),
    ("fecha_estado_actual", ["fecha estado actual"]),
    ("anio_estado_actual", ["ano de estado actual"]),
    ("estado_actual", ["estado actual", "estado"]),
]
DESTINO = [c for c, _ in COLUMNAS]
OBLIGATORIAS = ["municipio", "nombre", "dpi", "estado_actual"]


def norm_encabezado(s: object) -> str:
    s = "" if s is None else str(s)
    s = s.replace("﻿", "")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.lower().replace("ñ", "n")
    return re.sub(r"\s+", " ", s).strip()


class ErrorEstructura(ValueError):
    pass


def mapear(encabezados: list) -> tuple[list[int | None], list[str], list[str]]:
    """Devuelve, para cada columna destino, el índice en el archivo (o None);
    además la lista de columnas obligatorias faltantes y las columnas del archivo no reconocidas."""
    norm = [norm_encabezado(h) for h in encabezados]
    usados: set[int] = set()
    idx: list[int | None] = []
    for destino, alias in COLUMNAS:
        pos = None
        for a in alias:
            for i, h in enumerate(norm):
                if i not in usados and h == a:
                    pos = i
                    break
            if pos is not None:
                break
        if pos is not None:
            usados.add(pos)
        idx.append(pos)
    faltan = [d for d, i in zip(DESTINO, idx) if i is None and d in OBLIGATORIAS]
    extra = [encabezados[i] for i in range(len(encabezados)) if i not in usados and norm[i]]
    return idx, faltan, extra


def detectar_codificacion(muestra: bytes) -> str:
    try:
        muestra.decode("utf-8")
        return "utf-8-sig" if muestra.startswith(b"\xef\xbb\xbf") else "utf-8"
    except UnicodeDecodeError:
        pass
    buenos = set("áéíóúÁÉÍÓÚñÑüÜ")
    malos = set("„”“‚ƒ…†‡ˆ‰Š‹ŒŽ‘’•–—˜™š›œžŸ¿¡¬√≈∆«»")
    mejor, puntaje = "cp1252", -10**9
    for enc in ("mac_roman", "cp1252", "latin-1"):
        try:
            txt = muestra.decode(enc, errors="replace")
        except LookupError:
            continue
        p = sum(ch in buenos for ch in txt) * 2 - sum(ch in malos for ch in txt) * 3 - txt.count("�") * 5
        if p > puntaje:
            mejor, puntaje = enc, p
    return mejor


def _texto_celda(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    s = str(v).strip()
    return s if s != "" else None


class Lectura:
    """Iterador de filas ya mapeadas. Atributos: formato, codificacion, encabezados, faltan, extra."""

    def __init__(self, ruta: str | Path):
        self.ruta = Path(ruta)
        self.formato = "xlsx" if self.ruta.suffix.lower() in (".xlsx", ".xlsm") else "csv"
        self.codificacion: str | None = None
        self.separador: str | None = None
        self.encabezados: list = []
        self.faltan: list[str] = []
        self.extra: list[str] = []
        self.linea_encabezado = 1

    def __iter__(self) -> Iterator[tuple[int, list[str | None]]]:
        return self._csv() if self.formato == "csv" else self._xlsx()

    # --- CSV ---------------------------------------------------------
    def _csv(self):
        with open(self.ruta, "rb") as fb:
            muestra = fb.read(4_000_000)
        self.codificacion = detectar_codificacion(muestra)
        texto = muestra.decode(self.codificacion, errors="replace")
        primera = texto.splitlines()[0] if texto else ""
        self.separador = max([";", ",", "\t", "|"], key=primera.count)
        with open(self.ruta, "r", encoding=self.codificacion, errors="replace", newline="") as f:
            lector = csv.reader(f, delimiter=self.separador)
            filas_previas = []
            for n, fila in enumerate(lector, start=1):
                if not self.encabezados:
                    idx, faltan, extra = mapear(fila)
                    if not faltan:
                        self.encabezados, self.faltan, self.extra, self.linea_encabezado = fila, faltan, extra, n
                        continue
                    filas_previas.append(fila)
                    if n >= 15:
                        _, faltan, _ = mapear(filas_previas[0])
                        raise ErrorEstructura("No se encontró la fila de encabezados. Faltan columnas obligatorias: " + ", ".join(faltan))
                    continue
                if not any(c.strip() for c in fila):
                    continue
                yield n, [(fila[i].strip() or None) if i is not None and i < len(fila) else None for i in idx]
            if not self.encabezados:
                raise ErrorEstructura("El archivo no contiene encabezados reconocibles")

    # --- Excel -------------------------------------------------------
    def _xlsx(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.ruta, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                filas = ws.iter_rows(values_only=True)
                idx = None
                for n, fila in enumerate(filas, start=1):
                    if idx is None:
                        cand, faltan, extra = mapear(list(fila))
                        if not faltan:
                            idx = cand
                            self.encabezados, self.faltan, self.extra, self.linea_encabezado = list(fila), [], extra, n
                            self.codificacion = f"xlsx:{ws.title}"
                        elif n >= 15:
                            break
                        continue
                    valores = [_texto_celda(fila[i]) if i is not None and i < len(fila) else None for i in idx]
                    if not any(valores):
                        continue
                    yield n, valores
                if idx is not None:
                    return
            raise ErrorEstructura("Ninguna hoja tiene los encabezados obligatorios: " + ", ".join(OBLIGATORIAS))
        finally:
            wb.close()
